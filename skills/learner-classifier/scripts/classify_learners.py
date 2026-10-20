#!/usr/bin/env python3
"""
classify_learners.py — slow / advanced learner classification + reports.

Faithful port of the core logic in LearnerPDC (farm/worker/logic.py):
  * DataReader      -> read_marks / read_cgpa_map / read_grade_map
  * StudentDataProcessor -> process / filter_students
  * Format 1, 2, 3, 1+2 docx formatters
  * LibreOffice docx -> pdf conversion

Cryptographic PDF signing (endesive + PEM keys) is intentionally NOT ported:
it needs private key material, which belongs in the app's encrypted vault.

Usage:
  python classify_learners.py MARKS.xlsx --semester iii [options]
Run with -h for every option. Prints a JSON summary to stdout.
"""
import argparse
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

# --- CONFIGURATION (matches the app) -----------------------------------------
MIDTERM_TOTAL_MARKS = 30
DEFAULT_SLOW = 40.0       # strictly below  -> slow learner
DEFAULT_ADVANCED = 90.0   # strictly above  -> advanced learner
HEADER_SKIP_ROWS = 12     # data rows start after 12 header rows
COL_NAME, COL_REG, COL_TOTAL = 3, 4, 24   # columns D, E, Y (0-indexed)
MIN_COLUMNS = 25
PASSING_GRADES = {'A+', 'A', 'B', 'C', 'D', 'E', 'S'}
SEMESTER_MAPPING = {
    'i': 'I Year/ I semester', 'ii': 'I Year/ II semester', 'iii': 'II Year/ III semester',
    'iv': 'II Year/ IV semester', 'v': 'III Year/ V semester', 'vi': 'III Year/ VI semester',
    'vii': 'IV Year/ VII semester', 'viii': 'IV Year/ VIII semester',
}
HEADER_LINES = ['Manipal Institute of Technology', 'MAHE Manipal',
                'Computer Science and Engineering Department']
BODY_FONT = "Times New Roman"

K_NAME = 'Student Name'
K_REG = 'Register Number of the Student'
K_MARKS = 'Midterm Exam Marks (Out of 30)'
K_PCT = 'MidtermPercentage'
K_SUBJ = 'Subject Name'
K_SEM = 'Semester'
K_CGPA = 'CGPA (up to previous semester)'
K_ACT = 'Actions taken to improve performance'
K_REM = 'Remarks if any'
K_OUT = 'Outcome (Based on clearance in end-semester or makeup exam)'


# --- Reading ------------------------------------------------------------------
def _engine(path):
    return 'xlrd' if path.lower().endswith('.xls') else 'openpyxl'


def normalize_registration_number(reg_num) -> str:
    if reg_num is None or (not isinstance(reg_num, str) and pd.isna(reg_num)):
        return ''
    s = str(reg_num).strip().replace('\xa0', '').replace(' ', '').upper()
    s = re.sub(r'\.0+$', '', s)
    if '.' in s:
        try:
            s = str(int(float(s)))
        except (ValueError, TypeError):
            pass
    return s


def extract_subject(path):
    """Cell C2 holds '[CODE] Subject Name' -> 'Subject Name (CODE)'."""
    try:
        df = pd.read_excel(path, sheet_name=0, engine=_engine(path), nrows=2, header=None)
        val = df.iloc[1, 2] if df.shape[0] > 1 and df.shape[1] > 2 else None
        if val is None or pd.isna(val):
            return "Unknown Subject"
        val = str(val).strip()
        m = re.match(r'\[([A-Za-z0-9\s]+)\]\s*(.+)', val)
        if m:
            code, name = m.group(1).strip(), m.group(2).strip()
            return f"{name} ({code})" if code else name
        return val or "Unknown Subject"
    except Exception:
        return "Unknown Subject"


def read_marks(path):
    subject = extract_subject(path)
    try:
        df = pd.read_excel(path, sheet_name=0, engine=_engine(path), header=None,
                           skiprows=HEADER_SKIP_ROWS)
    except Exception as exc:
        raise ValueError(f"Could not read the marks file ({exc}).") from exc
    if df.shape[1] < MIN_COLUMNS:
        raise ValueError(
            f"Marks sheet has {df.shape[1]} column(s); the institutional layout needs at least "
            f"{MIN_COLUMNS} (through the Total(30) column Y). Is this the right sheet?")
    records = []
    for _, row in df.iterrows():
        reg = normalize_registration_number(row.iloc[COL_REG])
        if not reg or reg == 'NAN':
            continue
        name = str(row.iloc[COL_NAME]).strip() if pd.notna(row.iloc[COL_NAME]) else ''
        marks = row.iloc[COL_TOTAL] if pd.notna(row.iloc[COL_TOTAL]) else ''
        records.append({K_NAME: name, K_REG: reg, K_MARKS: marks})
    if not records:
        raise ValueError("No student rows found below the 12 header rows.")
    return records, subject


def _read_table(path, first_visible_sheet=False):
    if path.lower().endswith('.csv'):
        return pd.read_csv(path)
    sheet = 0
    if first_visible_sheet and _engine(path) == 'openpyxl':
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            vis = [ws.title for ws in wb.worksheets if ws.sheet_state == 'visible']
            wb.close()
            sheet = vis[0] if vis else 0
        except Exception:
            sheet = 0
    return pd.read_excel(path, sheet_name=sheet, engine=_engine(path))


def read_cgpa_map(path):
    """Col A = enrollment no., col B = CGPA."""
    if not path:
        return {}
    df = _read_table(path, first_visible_sheet=True)
    if len(df.columns) < 2:
        raise ValueError("CGPA file needs two columns: Enrollment No., Net Semester CGPA.")
    out = {}
    for _, row in df.iterrows():
        roll, val = normalize_registration_number(row.iloc[0]), row.iloc[1]
        if roll and pd.notna(val):
            try:
                out[roll] = f"{float(val):.2f}"
            except (ValueError, TypeError):
                out[roll] = str(val).strip()
    return out


def read_grade_map(path, course=None):
    """2 cols: reg, grade.  3+ cols: reg, course, grade (filtered to this course)."""
    if not path:
        return {}
    df = _read_table(path)
    if len(df.columns) < 2:
        raise ValueError("Grade file needs (Enrollment Id., Grade) or (Enrollment Id., Course Code, Grade).")
    out = {}
    if len(df.columns) >= 3:
        if course:
            code = course.split('(')[-1].replace(')', '').strip() if '(' in course else course
            name = course.split('(')[0].strip() if '(' in course else course
            col = df[df.columns[1]].astype(str).str.replace(' ', '').str.lower()
            mask = (col.str.contains(code.replace(' ', '').lower(), regex=False, na=False) |
                    col.str.contains(name.replace(' ', '').lower(), regex=False, na=False))
            if mask.any():
                df = df[mask]
        gi = 2
    else:
        gi = 1
    for _, row in df.iterrows():
        roll, g = normalize_registration_number(row.iloc[0]), row.iloc[gi]
        if roll and pd.notna(g):
            out[roll] = str(g).strip()
    return out


# --- Processing ---------------------------------------------------------------
def calc_pct(marks):
    try:
        return (float(marks) / MIDTERM_TOTAL_MARKS) * 100
    except Exception:
        return 0.0   # blank / 'AB' / non-numeric -> 0%, exactly like the app


def classify(pct, slow, adv):
    if pct < slow:
        return 'Slow'
    if pct > adv:
        return 'Advanced'
    return 'Average'


def process(records, subject, semester, actions, remarks, cgpa_map, grade_map, slow, adv):
    for s in records:
        s['Marks Numeric'] = _is_number(s[K_MARKS])
        s[K_PCT] = calc_pct(s[K_MARKS])
        s['Category'] = classify(s[K_PCT], slow, adv)
        s[K_SUBJ] = str(subject).strip()
        s[K_SEM] = str(semester).strip().lower()
        roll = normalize_registration_number(s[K_REG])
        s[K_CGPA] = cgpa_map.get(roll, '')
        s[K_ACT] = actions
        s[K_REM] = remarks
        g = str(grade_map.get(roll, '')).strip().upper()
        if not g or g in ('NAN', 'NONE'):
            s[K_OUT] = ''
        elif g in PASSING_GRADES:
            s[K_OUT] = 'Improved'
        else:
            s[K_OUT] = 'Not Improved'
    return records


def _is_number(v):
    try:
        float(v)
        return True
    except Exception:
        return False


def filter_students(students, learner_type, slow, adv):
    if learner_type == 'slow':
        f = [s for s in students if s[K_PCT] < slow]
    else:
        f = [s for s in students if s[K_PCT] > adv]
    return sorted(f, key=lambda s: s.get(K_REG, ''))


# --- Report formatting (python-docx) ------------------------------------------
class Formatter:
    def __init__(self, learner_type, subject, faculty_name=None, signature_image=None):
        self.learner_type = learner_type
        self.subject = subject
        self.faculty_name = faculty_name
        self.sig = signature_image

    @staticmethod
    def sem_string(roman):
        return SEMESTER_MAPPING.get(str(roman).strip().lower(), str(roman))

    @staticmethod
    def cell(cell, text, bold=False, size=10, align='LEFT', valign='TOP', font=None):
        cell.text = ''
        p = cell.add_paragraph()
        r = p.add_run(str(text))
        r.font.size = Pt(size)
        r.bold = bold
        if font:
            r.font.name = font
        p.alignment = getattr(WD_ALIGN_PARAGRAPH, align.upper(), WD_ALIGN_PARAGRAPH.LEFT)
        cell.vertical_alignment = getattr(WD_ALIGN_VERTICAL, valign.upper(), WD_ALIGN_VERTICAL.TOP)

    def signature(self, target):
        if self.sig:
            p = target.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            try:
                p.add_run().add_picture(io.BytesIO(self.sig), width=Inches(0.8))
            except Exception:
                pass
        p = target.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run("_" * 40 + "\n")
        if self.faculty_name:
            p.add_run(f"{self.faculty_name}\n")
        p.add_run("Signature of the\nsubject teacher / class coordinator")

    @staticmethod
    def header(cell):
        for line in HEADER_LINES:
            p = cell.add_paragraph()
            p.add_run(line).bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _mark(self, run, kind):
        if self.learner_type == kind:
            run.font.underline = True
            run.font.color.rgb = RGBColor(255, 0, 0)

    def format1(self, doc, s, slow, adv, page_break=False):
        h = doc.add_heading('Format 1. Assessment of the learning levels of the students:', level=2)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if page_break:
            h.paragraph_format.page_break_before = True
        ct = doc.add_table(rows=5, cols=1)
        ct.style = 'Table Grid'
        self.header(ct.cell(0, 0))
        st = ct.cell(1, 0).add_table(rows=4, cols=2)
        for i, (label, key) in enumerate([('Name of the Student:', K_NAME),
                                          ('Registration Number:', K_REG),
                                          ('Course:', K_SUBJ), ('Year /semester:', K_SEM)]):
            self.cell(st.cell(i, 0), label)
            v = s.get(key, '')
            if key == K_SEM:
                v = self.sem_string(v)
            elif key == K_SUBJ:
                v = str(v).upper()
            self.cell(st.cell(i, 1), v, font=BODY_FONT)
        pt = ct.cell(2, 0).add_table(rows=3, cols=4)
        pt.style = 'Table Grid'
        pt.cell(0, 2).merge(pt.cell(0, 3))
        for text, col in [('Sr. No.', 0), ('Parameter', 1), ('Weightage in Percentage', 2)]:
            self.cell(pt.cell(0, col), text, bold=True, align='CENTER')
        self.cell(pt.cell(1, 0), '1', align='CENTER')
        self.cell(pt.cell(1, 1), "Scores obtained by student class test / internal examination...\n"
                                 f"Considered Midterm exam conducted for {MIDTERM_TOTAL_MARKS}M:")
        self.cell(pt.cell(1, 2), f"{s.get(K_PCT, 0):.2f}", align='CENTER', font=BODY_FONT)
        self.cell(pt.cell(1, 3), "> %", align='CENTER')
        self.cell(pt.cell(2, 0), '2', align='CENTER')
        self.cell(pt.cell(2, 1), 'Performance of students in preceding university examination')
        self.cell(pt.cell(2, 2), str(s.get(K_CGPA, '')), align='CENTER', font=BODY_FONT)
        self.cell(pt.cell(2, 3), "> %", align='CENTER')
        ct.cell(3, 0).text = "Total Weightage"
        fc = ct.cell(4, 0)
        p1 = fc.add_paragraph()
        p1.add_run(f"1. Midterm score less than {_fmt_num(slow)}% considered as a ")
        self._mark(p1.add_run("slow learner"), 'slow')
        p2 = fc.add_paragraph()
        p2.add_run(f"2. Midterm score more than {_fmt_num(adv)}% considered as an ")
        self._mark(p2.add_run("advanced learner"), 'advanced')
        p2.add_run(" **")
        fc.add_paragraph().add_run(f"Date: {datetime.now():%d-%m-%Y}").font.name = BODY_FONT
        self.signature(fc)

    def format2(self, doc, s, page_break=False):
        h = doc.add_paragraph()
        h.style = 'Heading 2'
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if page_break:
            h.paragraph_format.page_break_before = True
        h.add_run('Format -2 Report of performance/ improvement for ')
        self._mark(h.add_run('slow'), 'slow')
        h.add_run(' and ')
        self._mark(h.add_run('advanced'), 'advanced')
        h.add_run(' learners')
        self.header(doc.add_table(rows=1, cols=1).cell(0, 0))
        ct = doc.add_table(rows=8, cols=2)
        ct.style = 'Table Grid'
        fields = [('1. Registration Number', K_REG), ('2. Name of the student', K_NAME),
                  ('3. Course', K_SUBJ), ('4. Year/Semester', K_SEM),
                  ('5. Midterm Percentage', K_PCT),
                  ('6. Activities/ Measure/special programs\ntaken to improve the performance', K_ACT),
                  ('7. Progress', K_OUT), ('Comments/remarks', K_REM)]
        for i, (label, key) in enumerate(fields):
            self.cell(ct.cell(i, 0), label)
            v = s.get(key, '')
            if key == K_SEM:
                v = self.sem_string(v)
            elif key == K_PCT:
                v = f"{v:.2f}%"
            elif key == K_SUBJ:
                v = str(v).upper()
            self.cell(ct.cell(i, 1), str(v).replace(';', '\n'), font=BODY_FONT)
        doc.add_paragraph().add_run(f"\nDate: {datetime.now():%d-%m-%Y}").font.name = BODY_FONT
        self.signature(doc)

    @staticmethod
    def _doc():
        d = Document()
        for sec in d.sections:
            sec.top_margin = Inches(0.5)
        return d

    def build(self, fmt, students, slow, adv):
        d = self._doc()
        if fmt == '1':
            for i, s in enumerate(students):
                self.format1(d, s, slow, adv, page_break=i > 0)
        elif fmt == '2':
            for i, s in enumerate(students):
                self.format2(d, s, page_break=i > 0)
        elif fmt == '4':
            for i, s in enumerate(students):
                self.format1(d, s, slow, adv, page_break=i > 0)
                self.format2(d, s, page_break=True)
        else:  # '3' summary
            self.summary(d, students)
        return d

    def summary(self, d, students):
        cols = ['Sl. No', 'Reg Number', 'Name of the student', 'Midterm Percentage', 'Progress']
        d.add_paragraph(f"Course: {str(self.subject).upper()}", style='Heading 3')
        t = d.add_table(rows=1, cols=len(cols))
        t.style = 'Table Grid'
        for j, c in enumerate(cols):
            self.cell(t.cell(0, j), c, bold=True)
        if not students:   # app prints an 8-row table with NIL in row 4
            for i in range(8):
                rc = t.add_row().cells
                self.cell(rc[0], str(i + 1), font=BODY_FONT)
                if i == 3:
                    self.cell(rc[2], "NIL", font=BODY_FONT, align='CENTER')
        else:
            for i, s in enumerate(students):
                rc = t.add_row().cells
                self.cell(rc[0], str(i + 1))
                self.cell(rc[1], s[K_REG])
                self.cell(rc[2], s[K_NAME])
                self.cell(rc[3], f"{s[K_PCT]:.2f}")
                self.cell(rc[4], s[K_OUT])
        self.signature(d)


def _fmt_num(x):
    return int(x) if float(x).is_integer() else x


# --- Writers -------------------------------------------------------------------
def soffice():
    for c in ('soffice', 'libreoffice'):
        if shutil.which(c):
            return c
    mac = '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    return mac if os.path.exists(mac) else None


def write_pdf(doc, out_path):
    exe = soffice()
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not found; cannot convert to PDF.")
    with tempfile.TemporaryDirectory() as td:
        tmp = os.path.join(td, "temp.docx")
        doc.save(tmp)
        subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', td, tmp],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        pdf = os.path.join(td, "temp.pdf")
        if not os.path.exists(pdf):
            raise RuntimeError("LibreOffice did not produce a PDF.")
        shutil.move(pdf, out_path)


def write_spreadsheet(students, path, slow, adv):
    cols = [K_REG, K_NAME, K_MARKS, K_PCT, 'Category', K_CGPA, K_OUT]
    df = pd.DataFrame(students).sort_values(K_REG)
    df[K_PCT] = df[K_PCT].round(2)
    df = df[cols].rename(columns={K_REG: 'Registration Number', K_MARKS: 'Midterm Marks (/30)',
                                  K_PCT: 'Midterm %', K_CGPA: 'CGPA', K_OUT: 'Progress'})
    if path.lower().endswith('.csv'):
        df.to_csv(path, index=False)
        return
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        df.to_excel(xw, sheet_name='All Students', index=False)
        df[df['Category'] == 'Slow'].to_excel(xw, sheet_name=f'Slow (<{_fmt_num(slow)}%)', index=False)
        df[df['Category'] == 'Advanced'].to_excel(xw, sheet_name=f'Advanced (>{_fmt_num(adv)}%)', index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                width = max(len(str(c.value or '')) for c in col) + 2
                ws.column_dimensions[col[0].column_letter].width = min(width, 45)
            ws.freeze_panes = 'A2'


# --- Main ---------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('marks', help='Institutional mid-term marks sheet (.xlsx/.xls)')
    ap.add_argument('--semester', required=True, help='Roman numeral i..viii')
    ap.add_argument('--cgpa', help='Optional CGPA file (.xlsx/.xls/.csv)')
    ap.add_argument('--grades', help='Optional end-sem grade file (.xlsx/.xls/.csv)')
    ap.add_argument('--slow', type=float, default=DEFAULT_SLOW, help='Slow threshold %% (default 40, strict <)')
    ap.add_argument('--advanced', type=float, default=DEFAULT_ADVANCED, help='Advanced threshold %% (default 90, strict >)')
    ap.add_argument('--learner-type', choices=['slow', 'advanced', 'both'], default='both')
    ap.add_argument('--format', choices=['1', '2', '3', '4', '5', 'none'], default='5',
                    help='1=Format1, 2=Format2, 3=Summary, 4=Format1+2 combined, 5=Combined+Summary, none=no reports')
    ap.add_argument('--output', choices=['word', 'pdf', 'both'], default='both')
    ap.add_argument('--actions', default='', help='"Actions taken to improve performance" text (; = new line)')
    ap.add_argument('--remarks', default='', help='Common remarks text')
    ap.add_argument('--faculty-name')
    ap.add_argument('--signature-image', help='PNG/JPG embedded above the signature line')
    ap.add_argument('--sheet-format', choices=['xlsx', 'csv'], default='xlsx')
    ap.add_argument('--outdir', default='learner_reports')
    a = ap.parse_args()

    if not 0 <= a.slow <= 100 or not 0 <= a.advanced <= 100:
        sys.exit("Thresholds must be between 0 and 100.")
    if a.slow > a.advanced:
        sys.exit("Slow threshold cannot exceed the advanced threshold.")

    records, subject = read_marks(a.marks)
    cgpa = read_cgpa_map(a.cgpa)
    grades = read_grade_map(a.grades, course=subject)
    students = process(records, subject, a.semester, a.actions, a.remarks, cgpa, grades, a.slow, a.advanced)

    os.makedirs(a.outdir, exist_ok=True)
    ds = datetime.now().strftime('%d_%m_%y')
    sn = a.semester.strip().upper()
    sub = re.sub(r'[\\/*?:"<>|()]', "", subject.replace(' ', '_'))
    sheet = os.path.join(a.outdir, f'{sub}_{sn}_Classified_{ds}.{a.sheet_format}')
    write_spreadsheet(students, sheet, a.slow, a.advanced)

    sig = open(a.signature_image, 'rb').read() if a.signature_image else None
    exts = {'word': ['docx'], 'pdf': ['pdf'], 'both': ['docx', 'pdf']}[a.output]
    types = ['slow', 'advanced'] if a.learner_type == 'both' else [a.learner_type]
    reports, groups = [], {}
    for lt in types:
        chosen = filter_students(students, lt, a.slow, a.advanced)
        groups[lt] = chosen
        if a.format == 'none':
            continue
        fmt = Formatter(lt, subject, a.faculty_name, sig)
        if not chosen:
            plan = [('3', 'Empty_Summary')]      # app falls back to an empty summary
        elif a.format == '5':
            plan = [('4', 'Combined'), ('3', 'Summary')]
        else:
            plan = [(a.format, {'1': 'Format1', '2': 'Format2', '3': 'Summary', '4': 'Combined'}[a.format])]
        for f, lbl in plan:
            doc = fmt.build(f, chosen, a.slow, a.advanced)
            for ext in exts:
                p = os.path.join(a.outdir, f'{sub}_{sn}_{lt.title()}Learner_{lbl}_{ds}.{ext}')
                if ext == 'docx':
                    doc.save(p)
                else:
                    write_pdf(doc, p)
                reports.append(p)

    non_numeric = [s for s in students if not s['Marks Numeric']]
    cats = pd.Series([s['Category'] for s in students]).value_counts().to_dict()
    summary = {
        'subject': subject,
        'semester': SEMESTER_MAPPING.get(a.semester.lower(), a.semester),
        'thresholds': {'slow_below': a.slow, 'advanced_above': a.advanced},
        'total_students': len(students),
        'counts': {k: int(cats.get(k, 0)) for k in ('Slow', 'Average', 'Advanced')},
        'slow_learners': [{'reg': s[K_REG], 'name': s[K_NAME], 'pct': round(s[K_PCT], 2)}
                          for s in filter_students(students, 'slow', a.slow, a.advanced)],
        'advanced_learners': [{'reg': s[K_REG], 'name': s[K_NAME], 'pct': round(s[K_PCT], 2)}
                              for s in filter_students(students, 'advanced', a.slow, a.advanced)],
        'non_numeric_marks': [{'reg': s[K_REG], 'name': s[K_NAME], 'raw': str(s[K_MARKS])}
                              for s in non_numeric],
        'cgpa_matched': sum(1 for s in students if s[K_CGPA]) if a.cgpa else None,
        'grades_matched': sum(1 for s in students if s[K_OUT]) if a.grades else None,
        'spreadsheet': sheet,
        'reports': reports,
    }
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
